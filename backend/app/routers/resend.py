"""Resend Campaign endpoints: intentional resend to previously contacted recipients.

Integrates with the existing campaign sender pipeline. Creates campaigns that the
scheduler picks up and processes according to sending mode and daily limits.
"""
from __future__ import annotations

import csv
import json
import logging
from io import StringIO, BytesIO

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import get_current_user
from ..models import User, Outreach, Campaign, CampaignTarget, CampaignStatus, SendMode, TargetState
from ..services import campaigns as cm

log = logging.getLogger("bevigrow.resend")

router = APIRouter(prefix="/api/resend", tags=["resend"])


@router.get("/health")
def resend_health():
    """Health check for resend router."""
    return {"status": "ok", "router": "resend"}


def parse_file(file: UploadFile) -> list[dict]:
    """Parse CSV or XLSX file and return list of rows."""
    log.info(f"parse_file called for: {file.filename if file else 'None'}")
    try:
        if not file or not file.filename:
            raise ValueError("No file provided")

        log.info(f"Reading file content from {file.filename}")
        content = file.file.read()
        log.info(f"Read {len(content)} bytes from file")

        if not content:
            raise ValueError("File is empty")

        filename = file.filename.lower()

        if filename.endswith(('.xlsx', '.xls')):
            try:
                wb = load_workbook(BytesIO(content))
                ws = wb.active
                if not ws:
                    raise ValueError("Workbook has no sheets")

                headers = None
                rows = []
                for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                    if row_idx == 1:
                        headers = [str(h).lower().strip() if h else f'col_{i}' for i, h in enumerate(row)]
                        continue
                    if not any(row):
                        continue
                    rows.append(dict(zip(headers or [], row)))
                return rows
            except Exception as e:
                log.error(f"Excel parsing error: {e}", exc_info=True)
                raise ValueError(f"Excel file error: {str(e)}")
        else:
            try:
                text_content = content.decode('utf-8', errors='replace')
                reader = csv.DictReader(StringIO(text_content))
                if not reader.fieldnames:
                    raise ValueError("CSV has no headers")
                rows = [row for row in reader if row and any(row.values())]
                return rows
            except Exception as e:
                log.error(f"CSV parsing error: {e}", exc_info=True)
                raise ValueError(f"CSV file error: {str(e)}")
    except ValueError:
        raise
    except Exception as e:
        log.error(f"File parsing error: {e}", exc_info=True)
        raise ValueError(f"File parsing error: {str(e)}")


def extract_email_field(row: dict) -> str | None:
    """Extract email from row, trying common field names."""
    email_fields = ['email', 'e-mail', 'mail', 'recipient_email', 'contact_email', 'email_address']
    for field in email_fields:
        value = row.get(field)
        if value:
            return str(value).strip().lower()
    return None


def check_contact_history(db: Session, email: str) -> dict:
    """Check if email was previously contacted."""
    try:
        stmt = select(Outreach).where(Outreach.email == email.lower()).order_by(Outreach.contacted_on.desc())
        outreach = db.scalars(stmt).first()

        return {
            'email': email,
            'is_in_history': outreach is not None,
            'last_sent_date': str(outreach.contacted_on) if outreach and outreach.contacted_on else None,
        }
    except Exception as e:
        log.warning(f"Contact history check failed for {email}: {e}")
        return {'email': email, 'is_in_history': False, 'last_sent_date': None}


@router.post("/review")
def resend_review(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    """Analyze uploaded file and identify previously contacted recipients."""
    log.info(f"[REVIEW] 1. Endpoint called for: {file.filename if file else 'unknown'}")

    try:
        log.info(f"[REVIEW] 2. File object: {file}")
        log.info(f"[REVIEW] 3. DB session: {db}")

        log.info(f"[REVIEW] 4. Starting _resend_review_impl")
        result = _resend_review_impl(file, db)

        log.info(f"[REVIEW] 5. Got result type: {type(result)}")
        log.info(f"[REVIEW] 6. Result keys: {result.keys() if isinstance(result, dict) else 'not a dict'}")
        log.info(f"[REVIEW] 7. Returning result")
        return result

    except HTTPException as e:
        log.error(f"[REVIEW] HTTPException: {e.status_code} - {e.detail}")
        raise
    except Exception as e:
        log.error(f"[REVIEW] EXCEPTION: {type(e).__name__}: {e}", exc_info=True)
        import traceback
        log.error(f"[REVIEW] TRACEBACK:\n{traceback.format_exc()}")

        return {
            "error": f"Server error: {type(e).__name__}",
            "detail": str(e)[:300],
            "type": type(e).__name__,
        }


def _resend_review_impl(file: UploadFile, db: Session):
    """Implementation of resend review logic."""
    log.info(f"[IMPL] 1. Started for file: {file.filename}")

    # Parse file
    try:
        log.info(f"[IMPL] 2. Calling parse_file")
        rows = parse_file(file)
        log.info(f"[IMPL] 3. Parsed {len(rows)} rows from file")
    except ValueError as e:
        log.warning(f"[IMPL] File parse error: {e}")
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        log.error(f"[IMPL] Unexpected error in parse_file: {type(e).__name__}: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"File parse error: {str(e)[:100]}")

    if not rows:
        log.info("[IMPL] 4. No rows found, returning empty")
        return {
            'total_recipients': 0,
            'previously_contacted': 0,
            'recipients': [],
        }

    log.info(f"[IMPL] 5. Processing {len(rows)} rows")
    recipients = []
    previously_contacted = 0

    # Process each row
    for idx, row in enumerate(rows):
        try:
            email = extract_email_field(row)
            if not email:
                log.debug(f"Row {idx}: No email field found")
                continue

            company = (row.get('company') or row.get('company_name') or 'Unknown')
            contact_person = row.get('contact_person') or row.get('contact')

            # Check history
            try:
                history = check_contact_history(db, email)
                log.debug(f"Email {email}: is_in_history={history['is_in_history']}")
            except Exception as history_error:
                log.error(f"Failed to check history for {email}: {history_error}")
                history = {'email': email, 'is_in_history': False, 'last_sent_date': None}

            recipient = {
                'email': email,
                'company_name': str(company) if company else 'Unknown',
                'contact_person': str(contact_person) if contact_person else None,
                'is_in_history': history['is_in_history'],
                'last_sent_date': history['last_sent_date'],
            }
            recipients.append(recipient)

            if history['is_in_history']:
                previously_contacted += 1
        except Exception as row_error:
            log.warning(f"Row {idx} processing error: {row_error}", exc_info=True)
            continue

    log.info(f"Review complete: {len(recipients)} total, {previously_contacted} previously contacted")
    result = {
        'total_recipients': len(recipients),
        'previously_contacted': previously_contacted,
        'recipients': recipients,
    }
    log.debug(f"Returning review result: {result}")
    return result


@router.post("/send")
def resend_send(
    file: UploadFile = File(...),
    campaign_name: str = Form(...),
    sending_mode: str = Form(...),
    resend_reason: str = Form(...),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Queue resend for previously contacted recipients."""
    log.info(f"Queuing resend: campaign={campaign_name}, mode={sending_mode}, reason={resend_reason}, user={user.email}")

    # Safety net: guarantee we always return valid JSON or HTTPException
    try:
        return _resend_send_impl(file, campaign_name, sending_mode, resend_reason, db, user)
    except HTTPException:
        raise
    except Exception as e:
        log.error(f"CRITICAL: Unexpected exception in resend_send: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Server error: {str(e)[:100]}")


def _resend_send_impl(file: UploadFile, campaign_name: str, sending_mode: str, resend_reason: str, db: Session, user):
    """Create a resend campaign and queue targets with existing sender pipeline."""
    try:
        rows = parse_file(file)
        log.info(f"Successfully parsed {len(rows)} rows for send")
    except ValueError as e:
        log.warning(f"File parse error: {e}")
        raise HTTPException(status_code=400, detail=str(e))

    if not rows:
        log.info(f"No rows in file for campaign '{campaign_name}'")
        raise HTTPException(status_code=400, detail="File contains no contacts to resend.")

    # Collect resend targets (only previously contacted)
    resend_targets = []
    position = 0

    for idx, row in enumerate(rows):
        try:
            email = extract_email_field(row)
            if not email:
                log.debug(f"Row {idx}: No email field, skipping")
                continue

            # Check if email was previously contacted
            try:
                history = check_contact_history(db, email)
            except Exception as history_error:
                log.error(f"Failed to check history for {email}: {history_error}")
                continue

            if not history['is_in_history']:
                log.debug(f"Email {email}: not in history, skipping")
                continue

            # This is a target to resend to
            company = (row.get('company') or row.get('company_name') or 'Unknown')
            contact_person = row.get('contact_person') or row.get('contact')

            resend_targets.append({
                'position': position,
                'email': email,
                'company_name': str(company) if company else 'Unknown',
                'contact_person': str(contact_person) if contact_person else None,
                'resend_reason': resend_reason[:255] if resend_reason else None,  # Truncate to DB limit
            })
            position += 1
            log.debug(f"Email {email}: added to resend queue")

        except Exception as e:
            log.warning(f"Row {idx} processing error: {e}", exc_info=True)
            continue

    if not resend_targets:
        log.warning(f"No previously contacted recipients found for resend campaign '{campaign_name}'")
        raise HTTPException(
            status_code=400,
            detail="No previously contacted recipients found in this file."
        )

    # Create Campaign record (integrates with existing sender pipeline)
    campaign = Campaign(
        name=campaign_name.strip() or "Untitled resend",
        mode=SendMode.automatic if sending_mode == "automatic" else SendMode.manual,
        daily_limit=min(50, cm.HARD_DAILY_CAP),  # Resend respects same 50/day limit
        owner_id=user.id,
        status=CampaignStatus.draft,
        source_filename=file.filename,
    )
    db.add(campaign)
    db.flush()
    log.info(f"Created resend campaign ID {campaign.id}: '{campaign_name}'")

    # Create CampaignTarget rows for each resend recipient
    targets = [
        CampaignTarget(
            campaign_id=campaign.id,
            position=target['position'],
            company_name=target['company_name'],
            email=target['email'],
            contact_person=target['contact_person'],
            state=TargetState.pending,
            resend_reason=target['resend_reason'],
        )
        for target in resend_targets
    ]
    db.add_all(targets)

    # Record the resend event
    cm.record(
        db,
        campaign.id,
        "resend_created",
        f"Resend campaign for {len(resend_targets)} previously contacted recipients. Reason: {resend_reason}",
    )

    db.commit()
    log.info(f"Queued {len(resend_targets)} targets for resend campaign '{campaign_name}'")

    return {
        'campaign_id': campaign.id,
        'campaign_name': campaign_name,
        'queued_count': len(resend_targets),
        'sending_mode': sending_mode,
        'resend_reason': resend_reason,
        'status': 'draft',
    }
